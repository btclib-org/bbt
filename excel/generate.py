# Copyright (c) The btclib developers
# Distributed under the MIT software license, see the accompanying
# LICENSE file or https://opensource.org/license/mit for the full text.

"""Write every workbook of this directory from the parameters below.

A workbook is not typed in Excel or Calc and then committed: it is built by
this script from the curve (or field) parameters given here as plain data,
so that a change to a workbook is a change to text and reviewed as one. Run
it with `uv run --locked python excel/generate.py` from a checkout; it
overwrites every `.xlsx` its own `CURVES` and `FINITE_FIELD_PRIMES` name.
`.github/scripts/check_generated_workbooks.py` is the gate that a committed
`.xlsx` still is this script's own output.

Every formula is a standard spreadsheet function -- `SUMPRODUCT`, `MOD`,
`ROW` -- and none is a macro, so the workbook this writes opens in Excel
or LibreOffice with no security prompt and no `#NAME?`. `inv_formula` and
`sqrt_formula` below are the two building blocks every other formula in
this file composes: a modular inverse and a modular square root, each
done by counting rows rather than by inverting or extracting a root
directly, since neither operation is a built-in spreadsheet function.

Excel and LibreOffice both recalculate a formula this script writes rather
than trusting a cached value, but for different reasons. `wb.calculation`
carries `fullCalcOnLoad="1"`, openpyxl's own default, which is Excel's own
documented instruction to recalculate the whole workbook on open regardless
of any cached result. openpyxl caches no result at all -- every formula
cell it writes carries an empty `<v/>` -- and LibreOffice's own import code
(`sc/source/filter/oox/formulabuffer.cxx`, `applySharedFormulas`) marks
exactly such a cell dirty and recalculates it "even if AutoCalc is
disabled", which does not depend on
`Tools > Options > Calc > Formula > Recalculation on File Load` (schema
default `OOXMLRecalcMode = 1`, "never"): that setting decides whether a
*present* cached value is trusted, and there is none here to trust.
"""

from __future__ import annotations

import math
from dataclasses import dataclass
from pathlib import Path
from typing import Literal

import openpyxl
from openpyxl.chart import ScatterChart, Series
from openpyxl.chart.marker import Marker
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.worksheet import Worksheet

HERE = Path(__file__).resolve().parent

# The margin `ROW($A$1:$A$<prime + this>)` carries above `prime` in every
# inv_formula/sqrt_formula call below: comfortably more than the one row the
# formula needs past `prime` itself, and shared by every curve here rather
# than tuned per workbook.
ROW_HEADROOM = 50

# EMU per point, the unit openpyxl's LineProperties.w wants: the OOXML
# drawingml unit is 1/914400 inch, a point is 1/72 inch, so a point is
# 914400/72 of them.
EMU_PER_POINT = 12700

# The positive (odd-root, or +sqrt over R) branch is red and the negative
# (even-root, or -sqrt over R) branch is blue, on every chart of an "EC"
# sheet that draws both -- one choice, held to consistently rather than
# argued for.
POSITIVE_BRANCH_COLOR = "FF0000"
NEGATIVE_BRANCH_COLOR = "0000FF"

# Neutral, on every chart that draws the curve's own symmetry line: it is
# a reference the eye checks the two branches against, not a third branch
# of its own, so it neither competes with red nor with blue.
SYMMETRY_AXIS_COLOR = "808080"

# Dark and thin, on both of every chart's own axis lines: dark enough to
# read as the plot's own frame rather than as a fourth series, thin
# enough not to compete with what is plotted against it.
AXIS_LINE_COLOR = "000000"
AXIS_LINE_WIDTH_PT = 0.75

# The number of points a real-curve sweep samples, and the least this
# generator ever extends a domain that is only one root wide: both are
# choices about how the curve looks, not about what it is, so a curve
# with a wider real domain than this is still sampled at this same
# point count, spread thinner, rather than the count growing with it.
REAL_CURVE_POINTS = 121
REAL_CURVE_MIN_SPAN = 4.0

# `leftmost`, in _add_real_curve below, is `_real_roots`' own float answer,
# and x^3+ax+b there is exactly 0 only in exact arithmetic: floating-point
# rounding leaves it a magnitude like 1E-7 either side of 0, so the sample
# placed exactly at a root reads slightly negative as often as not. This is
# comfortably above that noise and comfortably below any real feature of a
# curve this file draws, every one spanning at least REAL_CURVE_MIN_SPAN.
CUBIC_ZERO_TOLERANCE = 1e-6

# `LineProperties.prstDash`'s own preset names -- named here rather than
# read off `openpyxl` at type-checking time, since the descriptor itself
# validates a written value at run time regardless of what a caller's own
# annotation says, and this is the same set the constructor accepts.
DashStyle = Literal[
    "solid", "dot", "dash", "lgDash", "dashDot", "lgDashDot", "lgDashDotDot",
    "sysDash", "sysDot", "sysDashDot", "sysDashDotDot",
]


@dataclass(frozen=True)
class CurveSpec:
    """y^2 = x^3 + a*x + b over F(prime), generator G = (gx, gy) of order n.

    The filename is the tree's own naming pattern, `excel/README.md`'s:
    EC(a,b)-Fprime-G(gx,gy)-NX.xlsx, X the cardinality of the EC(Fp) group.
    That cardinality is `n * cofactor` and not `n` alone whenever G does
    not generate the whole group -- `n` here is the order of G itself, the
    number the "Finite Field" sheet's own table of multiples needs to stop
    at before it wraps around a second time.
    """

    a: int
    b: int
    prime: int
    gx: int
    gy: int
    n: int
    cofactor: int

    @property
    def filename(self) -> str:
        """The tree's own naming pattern for this curve's workbook."""
        cardinality = self.n * self.cofactor
        return (
            f"EC({self.a},{self.b})-F{self.prime}-G({self.gx},{self.gy})"
            f"-N{cardinality}.xlsx"
        )


CURVES = [
    # btclib-org/bbt#93: y^2 = x^3 + 2x + 5 over F101 has zero discriminant
    # and tabulates a node, not a group. This replacement keeps p = 101 and
    # is checked against btclib 2026.9.24: Curve(101, 7, 4, (0, 2), 97, 1,
    # False) is accepted, a brute-force count over F101 gives #E = 97, and
    # 97 is the smallest k with k*G at infinity -- so n below is measured
    # and not carried over from the curve it replaces.
    CurveSpec(a=7, b=4, prime=101, gx=0, gy=2, n=97, cofactor=1),
    # The other six curves this course also teaches, every parameter as
    # their own tables give it. `btclib`'s `Curve` is not what checks n
    # and cofactor here: its constructor refuses a composite n outright,
    # and 10, 270 and 280 below each factor beyond themselves and 1,
    # which a cryptographic curve never does but a teaching one may. n
    # is measured instead by a from-scratch Python double-and-add, as
    # the smallest positive k with k*G at infinity, and cofactor by
    # dividing that into a brute-force count of #E(Fp) over every x in
    # [0, prime).
    #
    # EC(-1,1)-F79-G(0,1) is the one curve here where G does not generate
    # the whole group: order(G) = 43 while #E(Fp) = 86, so cofactor = 2,
    # and the filename's own N is `n * cofactor` = 86 for exactly that
    # reason. Every other curve here has cofactor 1, where the two
    # coincide.
    CurveSpec(a=-1, b=1, prime=79, gx=0, gy=1, n=43, cofactor=2),
    CurveSpec(a=-7, b=10, prime=263, gx=3, gy=4, n=280, cofactor=1),
    CurveSpec(a=1, b=6, prime=11, gx=5, gy=9, n=13, cofactor=1),
    CurveSpec(a=2, b=3, prime=263, gx=200, gy=39, n=270, cofactor=1),
    CurveSpec(a=2, b=4, prime=7, gx=0, gy=2, n=10, cofactor=1),
    CurveSpec(a=6, b=9, prime=263, gx=0, gy=3, n=269, cofactor=1),
]

# The sheets of FiniteFields.xlsx, one per prime: F3, F5, F7, F11 and F79.
FINITE_FIELD_PRIMES = [3, 5, 7, 11, 79]
FINITE_FIELDS_FILENAME = "FiniteFields.xlsx"


def _real_roots(p: float, q: float) -> list[float]:
    """Real roots of x^3 + p*x + q = 0, ascending.

    The curve carries no x^2 term, so `p` and `q` are `a` and `b`
    themselves, needing no depression first. One real root or three --
    the sign of the discriminant decides which, and each case has its
    own closed form: Cardano's for one root, the trigonometric form for
    three. Every curve `CURVES` names takes the one-root branch, which
    `check_generated_workbooks.py` therefore exercises on every run;
    no curve here has three real roots, so the trigonometric branch is
    exercised by nothing in this tree.
    """
    disc = -4 * p**3 - 27 * q**2
    if disc > 0:
        r = 2 * math.sqrt(-p / 3)
        theta = math.acos(3 * q / (p * r)) / 3
        return sorted(r * math.cos(theta - 2 * math.pi * k / 3) for k in range(3))
    if disc < 0:

        def cbrt(x: float) -> float:
            return math.copysign(abs(x) ** (1 / 3), x)

        term = math.sqrt((q / 2) ** 2 + (p / 3) ** 3)
        return [cbrt(-q / 2 + term) + cbrt(-q / 2 - term)]
    # disc == 0: a repeated root, which a non-singular curve never carries
    # over F(prime) -- measured by hand for every curve `CURVES` names,
    # this file calling no check of its own -- but is not thereby ruled
    # out over R, the two discriminants answering unrelated questions.
    if p == 0 and q == 0:
        return [0.0]
    return sorted([3 * q / p, -3 * q / (2 * p), -3 * q / (2 * p)])


def _mod(expr: str, modulus: str) -> str:
    """`expr` reduced mod `modulus`, defensively, before it is used again.

    A difference of two field elements written directly into inv_formula's
    or sqrt_formula's `ROW(...)^...` term is exactly as valid mod `modulus`
    whether or not this wraps it -- MOD's own sign convention already
    matches the mathematical one -- but a second reduction of an argument
    that is sometimes negative costs nothing and is cheaper than trusting
    that every caller's expression happens to stay non-negative.
    """
    return f"MOD({expr},{modulus})"


def inv_formula(x_expr: str, prime_name: str, prime: int) -> str:
    """Invert `x_expr` mod `prime_name`, or answer NA() if `x_expr` is 0.

    `k` with `x*k == 1 (mod prime)` is unique in `[1, prime)`, so summing
    `k` over exactly the rows where the product matches finds it without a
    single division: `SUMPRODUCT` over `ROW($A$1:$A$<ceiling>)` for a
    `ceiling` comfortably past `prime`. `_mod` wraps `x_expr` before it is
    multiplied, for the reason given at its own definition above.

    0 has no inverse in a field, and no element's inverse is ever 0 --
    `SUMPRODUCT` answering 0 is exactly this case and never a legitimate
    result, so it is wrapped in `NA()` unconditionally, which is also what
    lets a `COUNT` elsewhere in this file skip the cell outright.
    """
    ceiling = prime + ROW_HEADROOM
    row_range = f"ROW($A$1:$A${ceiling})"
    raw = (
        f"SUMPRODUCT((MOD({_mod(x_expr, prime_name)}*{row_range},{prime_name})=1)"
        f"*({row_range}<{prime_name})*{row_range})"
    )
    return f"=IF({raw}=0,NA(),{raw})"


def sqrt_formula(n_expr: str, prime_name: str, prime: int) -> str:
    """Take a square root of `n_expr` mod `prime_name`, or NA() if none exists.

    Of the two roots `{r, prime - r}` a nonzero quadratic residue has,
    exactly one is odd whenever `prime` is odd -- their sum is `prime`
    itself -- so filtering `ROW(...)` to the odd values below `prime` and
    summing the one row where its square matches picks that root
    deterministically. The one root the odd filter cannot reach is 0
    itself, since `ROW($A$1:...)` never yields 0: `MOD(n_expr,prime)=0` is
    checked directly and returns 0 without going through `SUMPRODUCT` at
    all, rather than relying on the coincidence that a `SUMPRODUCT`
    finding nothing also answers 0. Whatever is left over a real
    non-residue, `SUMPRODUCT` finds nothing and answers 0 too, which is
    wrapped in `NA()` since 0 is never itself a genuine odd root.
    """
    ceiling = prime + ROW_HEADROOM
    row_range = f"ROW($A$1:$A${ceiling})"
    n_mod = _mod(n_expr, prime_name)
    raw = (
        f"SUMPRODUCT((MOD({row_range}^2,{prime_name})={n_mod})"
        f"*({row_range}<{prime_name})*(MOD({row_range},2)=1)*{row_range})"
    )
    return f"=IF({n_mod}=0,0,IF({raw}=0,NA(),{raw}))"


@dataclass(frozen=True)
class SeriesSpec:
    """One scatter series: its own y range, name and drawing style.

    `color` is a hex RGB triple with no leading `#`, `openpyxl`'s own
    convention; `None` leaves the theme's own default. `line` draws a
    segment between consecutive points -- the shape a path wants, and a
    discrete cloud of points does not -- and `marker` a dot at each one;
    at least one of the two is what makes a series visible at all.
    `line_width_pt` is `None` for the theme's own default width, or an
    explicit width in points for a line thinner (or thicker) than that.
    `x_range` is `None` for the chart's own shared `x_range` (every
    other series here plots against the same x), or a range of its own
    for a series -- the symmetry line below -- whose own points do not
    live in that column. `dash` is `None` for a solid line, or one of
    `LineProperties`' own preset dash names.
    """

    y_range: str
    name: str
    color: str | None = None
    line: bool = False
    marker: bool = True
    smooth: bool = False
    line_width_pt: float | None = None
    x_range: str | None = None
    dash: DashStyle | None = None


def _add_scatter(  # noqa: PLR0913
    ws: Worksheet,
    anchor: str,
    title: str,
    x_range: str,
    series_specs: list[SeriesSpec],
    *,
    cross_at_zero: bool = False,
) -> None:
    """Draw a scatter chart of `x_range` against each of `series_specs`.

    Six parameters and not five: `cross_at_zero` is one curve's own
    (the real one's) request among the several this function draws for,
    and a keyword-only flag reads plainer at its one call site than a
    config object built solely to carry it would.

    `cross_at_zero` sets each axis' own `crosses` to `"autoZero"`
    explicitly -- Excel's own documented default for an unset `crosses`,
    so this changes nothing left at that default, but an explicit
    setting is exercised on every run rather than assumed of every
    reader's own copy of Excel or LibreOffice.
    """
    chart = ScatterChart()
    chart.title = title
    chart.style = 13
    chart.x_axis.title = "x"
    chart.y_axis.title = "y"
    if cross_at_zero:
        chart.x_axis.crosses = "autoZero"
        chart.y_axis.crosses = "autoZero"
    # Neither axis draws its gridlines, on every chart this function
    # builds: a plotted curve or a table's own polyline reads against the
    # axis lines alone, with nothing behind it competing for attention.
    chart.x_axis.majorGridlines = None
    chart.y_axis.majorGridlines = None
    # `delete`'s own OOXML default is true, an axis openpyxl never sets
    # one way or the other left hidden rather than shown -- both drawn,
    # each with a thin, dark line of its own (`spPr`, not the theme's
    # default), numeric labels kept off the low edge of the plot rather
    # than beside wherever the axis itself crosses (`tickLblPos`, which
    # matters once `cross_at_zero` puts that crossing in the plot's own
    # middle rather than at a corner of it), and major ticks pointing
    # out from the line rather than across the data.
    for axis in (chart.x_axis, chart.y_axis):
        axis.delete = False
        axis.spPr = GraphicalProperties(
            ln=LineProperties(
                solidFill=AXIS_LINE_COLOR, w=round(AXIS_LINE_WIDTH_PT * EMU_PER_POINT),
            ),
        )
        axis.tickLblPos = "low"
        axis.majorTickMark = "out"
    # Only the x and y axes frame the plot: no outline box around either
    # the chart area or the plot area, a solid white fill rather than
    # the theme's own on either (so the chart reads the same over a
    # dark background too), and no rounded corners on the chart area
    # either -- `roundedCorners` is a chart-space property of its own,
    # separate from the `spPr` border/fill either level draws by default
    # when neither is set.
    chart.roundedCorners = False
    chart.graphical_properties = GraphicalProperties(
        solidFill="FFFFFF", ln=LineProperties(noFill=True),
    )
    chart.plot_area.graphicalProperties = GraphicalProperties(
        solidFill="FFFFFF", ln=LineProperties(noFill=True),
    )
    # every series here is drawn alike, so the chart-level hint is read
    # off the first one; a reader of the chart's own XML then finds the
    # same choice in two places rather than in the per-series lines alone
    first = series_specs[0]
    if first.smooth:
        chart.scatterStyle = "smoothMarker" if first.marker else "smooth"
    elif first.line:
        chart.scatterStyle = "lineMarker" if first.marker else "line"
    else:
        chart.scatterStyle = "marker"
    for spec in series_specs:
        series = Series(
            openpyxl.chart.Reference(range_string=f"'{ws.title}'!{spec.y_range}"),
            openpyxl.chart.Reference(
                range_string=f"'{ws.title}'!{spec.x_range or x_range}",
            ),
            title=spec.name,
            title_from_data=False,
        )
        series.smooth = spec.smooth
        if spec.marker:
            series.marker = Marker(symbol="circle", size=5)
            if spec.color:
                series.marker.graphicalProperties = GraphicalProperties(
                    solidFill=spec.color,
                    ln=LineProperties(solidFill=spec.color),
                )
        else:
            series.marker = Marker(symbol="none")
        if spec.line:
            width = (
                round(spec.line_width_pt * EMU_PER_POINT)
                if spec.line_width_pt is not None
                else None
            )
            if spec.color or width is not None or spec.dash is not None:
                series.graphicalProperties = GraphicalProperties(
                    ln=LineProperties(
                        solidFill=spec.color, w=width, prstDash=spec.dash,
                    ),
                )
        else:
            series.graphicalProperties = GraphicalProperties(
                ln=LineProperties(noFill=True),
            )
        chart.series.append(series)
    ws.add_chart(chart, anchor)


def _set_column_widths(ws: Worksheet, widths: dict[str, float]) -> None:
    """Set each column letter in `widths` to its own explicit width.

    Every width below is one Excel's own AutoFit produced for that sheet's
    fixed header text -- "y = -sqrt(x^3+ax+b)", "x^3+a*x+b mod p" and the
    rest are the same string whichever curve is being built -- and for a
    data column, over the digit range this course's own curves stay in:
    every `prime` this file names is at most three digits (the largest
    is 263) and every `a`, `b`, `gx`, `gy` is at most three digits too
    (the largest is `gx=200`), so a width AutoFit chose for one curve's
    table fits every other curve's table the same way. A curve needing
    a wider prime would need these revisited.
    """
    for letter, width in widths.items():
        ws.column_dimensions[letter].width = width


def _add_real_curve(ws: Worksheet, spec: CurveSpec) -> None:
    """Fill the "EC over R" sheet: the curve's own shape over R, and its chart.

    One column sweeps x from the curve's own leftmost real root to a
    bound past its rightmost one, and two formula columns beside it --
    each of a, b, and neither hardcoded -- answer y and -y, or NA() where
    x^3+ax+b is negative and neither root is real. A curve with three
    real roots has two components, [r1, r2] and [r3, infinity), with a
    gap between them where the same NA() check already applies: nothing
    here treats that case specially, because the check already does.
    """
    roots = _real_roots(spec.a, spec.b)
    leftmost = min(roots)
    scale = max([1.0, *(abs(r) for r in roots)])
    span = max(REAL_CURVE_MIN_SPAN, 3 * scale)
    right_bound = max(roots) + span
    step = (right_bound - leftmost) / (REAL_CURVE_POINTS - 1)

    step_row = 1
    header_row = step_row + 1
    first_row = header_row + 1
    last_row = first_row + REAL_CURVE_POINTS - 1

    ws.cell(step_row, 1, "step")
    ws.cell(step_row, 2, step)
    ws.cell(header_row, 1, "x (real)")
    ws.cell(header_row, 2, "y = sqrt(x^3+ax+b)")
    ws.cell(header_row, 3, "y = -sqrt(x^3+ax+b)")
    for offset in range(REAL_CURVE_POINTS):
        row = first_row + offset
        if offset == 0:
            ws.cell(row, 1, leftmost)
        else:
            ws.cell(row, 1, f"=A{row - 1}+$B${step_row}")
        cubic = f"(A{row}^3+a*A{row}+b)"
        clamped = f"MAX(0,{cubic})"
        ws.cell(
            row, 2, f"=IF({cubic}<-{CUBIC_ZERO_TOLERANCE},NA(),SQRT({clamped}))",
        )
        ws.cell(
            row, 3, f"=IF({cubic}<-{CUBIC_ZERO_TOLERANCE},NA(),-SQRT({clamped}))",
        )

    # The table sits in A:C, no column to its left, so the chart -- one
    # empty column (D) further on, at the header's own row -- sits beside
    # its top rather than below its last row: `last_row` grows with
    # `REAL_CURVE_POINTS`, and a chart anchored to it would not stay put.
    _set_column_widths(ws, {"A": 12.6640625, "B": 14.5, "C": 15.0})
    _add_scatter(
        ws,
        f"E{header_row}",
        f"y^2 = x^3 + {spec.a}x + {spec.b} over R",
        f"$A${first_row}:$A${last_row}",
        [
            SeriesSpec(
                f"$B${first_row}:$B${last_row}",
                "y = +sqrt(x^3+ax+b)",
                color=POSITIVE_BRANCH_COLOR,
                line=True,
                marker=False,
                smooth=True,
            ),
            SeriesSpec(
                f"$C${first_row}:$C${last_row}",
                "y = -sqrt(x^3+ax+b)",
                color=NEGATIVE_BRANCH_COLOR,
                line=True,
                marker=False,
                smooth=True,
            ),
        ],
        cross_at_zero=True,
    )


def _active_worksheet(wb: openpyxl.Workbook) -> Worksheet:
    """Narrow `wb.active` from "sheet, chartsheet or none" to "sheet".

    A freshly built `Workbook()` always has exactly one worksheet, but
    `.active`'s own type admits a chartsheet or no sheet at all -- narrowed
    here once rather than at every one of the many cells that follow.
    """
    active = wb.active
    if not isinstance(active, Worksheet):
        msg = "a new Workbook always has one worksheet"
        raise TypeError(msg)
    return active


def _build_ec_over_r_sheet(wb: openpyxl.Workbook, spec: CurveSpec) -> None:
    """Add the "EC over R" sheet: the curve's own shape over R, and its chart.

    This sheet is separate from the finite-field table's own sheet so the
    two charts do not compete for space on one: the reference to `a` and
    `b` is unaffected by which sheet is first, a defined name being a
    workbook-level binding rather than a sheet-level one, reading the
    same wherever the formula that uses it sits.
    """
    ws = _active_worksheet(wb)
    ws.title = "EC over R"
    _add_real_curve(ws, spec)


def _build_ec_over_f_sheet(wb: openpyxl.Workbook, spec: CurveSpec) -> None:
    """Add the "EC over F" sheet: every point of the curve over F(prime)."""
    ec = wb.create_sheet("EC over F")

    ec["A1"], ec["B1"] = "a", spec.a
    ec["A2"], ec["B2"] = "b", spec.b
    ec["A3"], ec["B3"] = "prime", spec.prime
    ec["A4"], ec["B4"] = "cofactor", spec.cofactor
    wb.defined_names["a"] = DefinedName("a", attr_text="'EC over F'!$B$1")
    wb.defined_names["b"] = DefinedName("b", attr_text="'EC over F'!$B$2")
    wb.defined_names["prime"] = DefinedName("prime", attr_text="'EC over F'!$B$3")

    # "# points" sits at row 5, between the parameters above it and the
    # table below: a blank row 6 separates it from `table_header_row`,
    # the same gap the "Finite Field" sheet below keeps for the same
    # reason -- a label that names a formula's own result reads better
    # apart from the table that formula counts.
    table_header_row = 7
    first_data_row = table_header_row + 1
    last_data_row = first_data_row + spec.prime - 1
    ec.cell(5, 1, "# points")
    ec.cell(
        5,
        2,
        f"=COUNT(C{first_data_row}:C{last_data_row})"
        f"+COUNT(D{first_data_row}:D{last_data_row})+1",
    )
    ec.cell(table_header_row, 1, "x")
    ec.cell(table_header_row, 2, "x^3+a*x+b mod p")
    ec.cell(table_header_row, 3, "y")
    ec.cell(table_header_row, 4, "-y")
    for offset in range(spec.prime):
        row = first_data_row + offset
        x = offset
        ec.cell(row, 1, x)
        ec.cell(row, 2, f"=MOD(A{row}^3+a*A{row}+b,prime)")
        ec.cell(row, 3, sqrt_formula(f"B{row}", "prime", spec.prime))
        ec.cell(
            row,
            4,
            f"=IF(OR(NOT(ISNUMBER(C{row})),C{row}=0),NA(),MOD(prime-C{row},prime))",
        )
    # (x, y) and (x, prime-y) are symmetric about y = prime/2 for every
    # point this table carries, so a horizontal line at that height is
    # every such pair's own axis of symmetry. Its two points sit in F1:G2,
    # otherwise unused, rather than in a column the table itself needs;
    # both x and y are formulas so the line follows `prime` wherever it
    # is bound, rather than a value copied out of it once.
    ec.cell(1, 6, 0)
    ec.cell(1, 7, "=prime/2")
    ec.cell(2, 6, "=prime-1")
    ec.cell(2, 7, "=prime/2")
    _set_column_widths(ec, {"A": 7.5, "B": 14.1640625, "C": 5.0, "D": 5.0})
    _add_scatter(
        ec,
        "H5",
        f"y^2 = x^3 + {spec.a}x + {spec.b} over F{spec.prime}",
        f"$A${first_data_row}:$A${last_data_row}",
        [
            SeriesSpec(
                f"$C${first_data_row}:$C${last_data_row}",
                "y (odd root)",
                color=POSITIVE_BRANCH_COLOR,
            ),
            SeriesSpec(
                f"$D${first_data_row}:$D${last_data_row}",
                "-y (even root)",
                color=NEGATIVE_BRANCH_COLOR,
            ),
            SeriesSpec(
                "$G$1:$G$2",
                "symmetry axis",
                color=SYMMETRY_AXIS_COLOR,
                line=True,
                marker=False,
                dash="dash",
                line_width_pt=0.75,
                x_range="$F$1:$F$2",
            ),
        ],
    )


def _build_finite_field_sheet(wb: openpyxl.Workbook, spec: CurveSpec) -> None:
    """Add the "Finite Field" sheet: every multiple of G, 1G to n*G."""
    ff = wb.create_sheet("Finite Field")
    ff["A1"], ff["B1"] = "Gx", spec.gx
    ff["A2"], ff["B2"] = "Gy", spec.gy
    wb.defined_names["Gx"] = DefinedName("Gx", attr_text="'Finite Field'!$B$1")
    wb.defined_names["Gy"] = DefinedName("Gy", attr_text="'Finite Field'!$B$2")

    # "# points" sits at row 3, between Gx/Gy above it and the table
    # below, the same placement "EC over F" gives its own count, and for
    # the same reason.
    mult_header_row = 5
    first_mult_row = mult_header_row + 1
    last_mult_row = first_mult_row + spec.n - 1
    ff.cell(3, 1, "# points")
    ff.cell(3, 2, f"=COUNT(B{first_mult_row}:B{last_mult_row})+1")
    headers = ((1, "k"), (2, "x"), (3, "y"), (4, "y^2 mod p"), (5, "x^3+a*x+b mod p"))
    for col, label in headers:
        ff.cell(mult_header_row, col, label)
    for offset in range(spec.n):
        row = first_mult_row + offset
        k = offset + 1
        ff.cell(row, 1, k)
        if k == 1:
            ff.cell(row, 2, "=Gx")
            ff.cell(row, 3, "=Gy")
        else:
            prev = row - 1
            ff.cell(row, 2, f"=MOD(J{prev}*J{prev}-B{prev}-Gx,prime)")
            ff.cell(row, 3, f"=MOD(J{prev}*(Gx-B{row})-Gy,prime)")
        ff.cell(row, 4, f"=MOD(C{row}*C{row},prime)")
        ff.cell(row, 5, f"=MOD(B{row}^3+a*B{row}+b,prime)")
        if k < spec.n:
            # The slope that carries this row's point to the next one:
            # the tangent at G doubling it for k == 1, the chord through G
            # and this row's point otherwise. Both feed row + 1's B and C
            # above, and the last row needs neither: its own B and C are
            # themselves where k*G first returns to infinity, `spec.n`
            # being the order of G and not the curve's own cardinality
            # (`CurveSpec`'s own docstring) -- the chord through G and
            # (n-1)*G is vertical, (n-1)*G and G sharing one x, since
            # (n-1)*G = -G whenever n is G's own order, and NA() ends up
            # exactly what inv_formula answers that vertical chord's
            # zero denominator with.
            #
            # Columns F and G..J, not G and H..K: a single empty column
            # (F) separates this block from the table to its left, and
            # column K stays empty too, with nothing written past J.
            if k == 1:
                num, den = f"=MOD(3*B{row}*B{row}+a,prime)", f"=MOD(2*C{row},prime)"
            else:
                num, den = f"=MOD(Gy-C{row},prime)", f"=MOD(Gx-B{row},prime)"
            ff.cell(row, 7, num)
            ff.cell(row, 8, den)
            ff.cell(row, 9, inv_formula(f"H{row}", "prime", spec.prime))
            ff.cell(row, 10, f"=MOD(G{row}*I{row},prime)")
    # (x, y) and (x, prime-y) are symmetric about y = prime/2 for every
    # point this table carries, the same fact "EC over F" draws its own
    # line for. Its two points sit in D1:E2, otherwise unused, rather
    # than in a column the table itself needs; both x and y are formulas
    # so the line follows `prime` wherever it is bound.
    ff.cell(1, 4, 0)
    ff.cell(1, 5, "=prime/2")
    ff.cell(2, 4, "=prime-1")
    ff.cell(2, 5, "=prime/2")
    _set_column_widths(
        ff,
        {
            "A": 7.1640625,
            "B": 4.1640625,
            "C": 5.0,
            "D": 9.0,
            "E": 14.1640625,
            "F": 5.0,
            "G": 4.1640625,
            "H": 3.1640625,
            "I": 5.0,
            "J": 5.0,
        },
    )
    _add_scatter(
        ff,
        "L5",
        f"multiples of G = ({spec.gx},{spec.gy}) over F{spec.prime}",
        f"$B${first_mult_row}:$B${last_mult_row}",
        [
            SeriesSpec(
                f"$C${first_mult_row}:$C${last_mult_row}",
                "k*G, k = 1, 2, ...",
                line=True,
                # a very thin line, so that the path reads as a path
                # across a dense table rather than as a thick band over
                # its own markers
                line_width_pt=0.25,
            ),
            SeriesSpec(
                "$E$1:$E$2",
                "symmetry axis",
                color=SYMMETRY_AXIS_COLOR,
                line=True,
                marker=False,
                dash="dash",
                line_width_pt=0.75,
                x_range="$D$1:$D$2",
            ),
        ],
    )


def build_ec_workbook(spec: CurveSpec) -> openpyxl.Workbook:
    """Build the workbook `spec` describes: its points, and G's multiples."""
    wb = openpyxl.Workbook()
    _build_ec_over_r_sheet(wb, spec)
    _build_ec_over_f_sheet(wb, spec)
    _build_finite_field_sheet(wb, spec)
    return wb


def _build_prime_demo_sheet(ws: Worksheet, prime: int) -> None:
    """Fill one sheet of FiniteFields.xlsx: Fp's opposite, inverse and roots.

    `prime` is this sheet's own -- a defined name scoped to it with
    `ws.defined_names`, not to the workbook, which is btclib-org/bbt#134:
    `calc/FiniteFields.ods` gave every sheet's `prime` the same one
    workbook-wide binding, so F3 through F11 computed modulo F79's own 79
    instead of their own. A `localSheetId`-scoped `definedName` is what
    OOXML calls this, and openpyxl writes one wherever a name is assigned
    through a worksheet's own `defined_names` rather than the workbook's.
    """
    ws.cell(1, 1, prime)
    ws.cell(1, 2, "opposite")
    ws.cell(1, 3, "inverse")
    ws.cell(1, 4, "odd sqrt")
    ws.cell(1, 5, "even sqrt")
    ws.cell(1, 7, "checks")
    ws.defined_names["prime"] = DefinedName("prime", attr_text=f"'{ws.title}'!$A$1")
    for x in range(prime):
        row = x + 2
        ws.cell(row, 1, x)
        ws.cell(row, 2, f"=MOD(prime-A{row},prime)")
        ws.cell(row, 3, inv_formula(f"A{row}", "prime", prime))
        ws.cell(row, 4, sqrt_formula(f"A{row}", "prime", prime))
        ws.cell(row, 5, f"=MOD(prime-D{row},prime)")
        ws.cell(row, 7, f"=MOD($A{row}+B{row},prime)")
        ws.cell(row, 8, f"=MOD($A{row}*C{row},prime)")
        ws.cell(row, 9, f"=MOD(D{row}*D{row},prime)")
        ws.cell(row, 10, f"=MOD(E{row}*E{row},prime)")


def build_finite_fields_workbook() -> openpyxl.Workbook:
    """Build FiniteFields.xlsx: one sheet per `FINITE_FIELD_PRIMES` entry."""
    wb = openpyxl.Workbook()
    for index, prime in enumerate(FINITE_FIELD_PRIMES):
        ws = _active_worksheet(wb) if index == 0 else wb.create_sheet()
        ws.title = f"F{prime}"
        _build_prime_demo_sheet(ws, prime)
    return wb


def main() -> None:
    """Write every workbook `CURVES` and `FINITE_FIELD_PRIMES` name."""
    for spec in CURVES:
        wb = build_ec_workbook(spec)
        wb.save(HERE / spec.filename)
        print(spec.filename)
    wb = build_finite_fields_workbook()
    wb.save(HERE / FINITE_FIELDS_FILENAME)
    print(FINITE_FIELDS_FILENAME)


if __name__ == "__main__":
    main()
