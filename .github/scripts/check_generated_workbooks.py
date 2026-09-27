# Copyright (c) The btclib developers
# Distributed under the MIT software license, see the accompanying
# LICENSE file or https://opensource.org/license/mit for the full text.

"""Check that every committed `excel/*.xlsx` is `excel/generate.py`'s output.

btclib-org/bbt#136 asks for a check "beside the generator that the
committed files are its output" -- the tree carries no test suite of its
own kind for a course-material generator, and a spreadsheet is a binary
that a `git diff` says nothing about, so this is the whole of what
answers whether an edit made to a workbook by hand, or a `generate.py`
changed without regenerating, has drifted from what the script would
write.

The comparison is cell by cell, not byte by byte: two workbooks holding
the same values and formulas can differ in a relationship id, a shared
string table's own order, or a timestamp, none of which is a change to
the material. Every chart's title, every series and every axis, and
the chart-space's and plot-area's own border and fill, are compared
too -- the chart XML part itself, parsed with `openpyxl`'s own
`ChartSpace.from_tree`, entered directly rather than through the
convenience wiring a loaded worksheet's own `_charts` goes through
afterward, which drops what the plot area's own graphical properties
are. Every column's own explicit width and every chart's own anchor
cell are compared as well, each read off the loaded worksheet, which
does carry those two forward. Every workbook-scoped and sheet-scoped
defined name is compared too, both its own name and the cell it binds
to.

Everything it reads it names, and it fails where it reads nothing: a
gate that opened no file is silent in exactly the way a gate that found
no defect is.
"""

from __future__ import annotations

import importlib.util
import io
import re
import sys
import zipfile
from pathlib import Path
from typing import TYPE_CHECKING, cast

import openpyxl
from openpyxl.chart.chartspace import ChartSpace
from openpyxl.xml.functions import fromstring, tostring

if TYPE_CHECKING:
    from types import ModuleType

    from openpyxl.descriptors.serialisable import _SerialisableTreeElement
    from openpyxl.workbook.workbook import Workbook
    from openpyxl.worksheet.worksheet import Worksheet

ROOT = Path(__file__).resolve().parents[2]
EXCEL_DIR = ROOT / "excel"

_CHART_PART = re.compile(r"xl/charts/chart(\d+)\.xml")


def _load_generate() -> ModuleType:
    """Import `excel/generate.py` as a module, without `excel/` on `sys.path`.

    A plain `import generate` would need `excel/` on `sys.path`, which
    nothing else this tree runs does; loading it from its own file path
    keeps this script runnable from any working directory.

    Registering the module in `sys.modules` before `exec_module` runs it
    is not optional: `excel/generate.py`'s `CurveSpec` and `SeriesSpec`
    are `@dataclass(frozen=True)` under `from __future__ import
    annotations`, and CPython's own `dataclasses._is_type` resolves such
    a class's annotations by looking its module back up in `sys.modules`
    -- a module executed without ever being entered there is not found,
    and the lookup raises `AttributeError` on `None.__dict__` rather than
    answering `False` and moving on.
    """
    spec = importlib.util.spec_from_file_location(
        "generate", EXCEL_DIR / "generate.py",
    )
    if spec is None or spec.loader is None:
        msg = "excel/generate.py: could not be loaded as a module"
        raise ImportError(msg)
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def _roundtrip(wb: Workbook) -> tuple[Workbook, bytes]:
    """Save `wb` and load it back, shaped like a workbook read from disk.

    A cell or a chart just built in memory and one `openpyxl.load_workbook`
    parsed back out of a file are not always the same Python shape for the
    same content -- a colour is a plain string until it has been through
    the XML and back, a `SolidColorFillProperties` afterwards. Sending the
    freshly generated workbook through the same round trip as the
    committed one puts both sides through identical code before either is
    compared, so the comparison below is never between two different
    representations of one answer.

    The bytes saved are returned alongside the workbook loaded from them:
    `_chart_xml` reads a chart's own XML part directly from those bytes,
    the one property (`ChartSpace.chart.plotArea.spPr`) a loaded chart's
    own convenience wiring does not carry forward into `ws._charts`.
    """
    buffer = io.BytesIO()
    wb.save(buffer)
    data = buffer.getvalue()
    buffer.seek(0)
    return openpyxl.load_workbook(buffer), data


def _cells(ws: Worksheet) -> dict[str, object]:
    """Return every non-empty cell of `ws`, keyed by its own coordinate."""
    return {
        cell.coordinate: cell.value
        for row in ws.iter_rows()
        for cell in row
        if cell.value is not None
    }


def _chart_parts(data: bytes) -> list[bytes]:
    """Return every `xl/charts/chartN.xml` part of a saved `.xlsx`, in order.

    Numeric order is creation order, `generate.py`'s own: every sheet in
    turn, one chart per sheet, so it lines up with each sheet's own
    `ws._charts` -- good for a count and an order, just not for what a
    loaded chart's own `.plot_area` carries (`_chart_xml`'s docstring).
    """
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        names = sorted(
            (n for n in archive.namelist() if _CHART_PART.fullmatch(n)),
            key=lambda n: int(_CHART_PART.fullmatch(n).group(1)),  # type: ignore[union-attr]
        )
        return [archive.read(name) for name in names]


def _chart_xml(data: bytes) -> bytes:
    """Serialise a chart XML part the way `openpyxl` itself parses it back.

    `ChartSpace.from_tree` is `openpyxl`'s own declarative parser for a
    chart part -- the whole `chartSpace`: plot area, every axis (each
    with its own gridlines, crossing point and scaling), legend and
    title alike, and the chart-space's and plot-area's own graphical
    properties. Entered here directly on the chart's own XML part, it
    answers something a loaded worksheet's own `_charts` does not: that
    convenience wiring reassigns a loaded chart's `x_axis`, `y_axis` and
    `series` from the parsed tree, but leaves its `.plot_area` the blank
    one `ChartBase.__init__` always builds, the parsed `<c:plotArea>`'s
    own `spPr` -- a border-and-fill removal there -- never copied onto
    it. Comparing what `ChartSpace.from_tree` itself answers is what
    recovers that, invisible to any comparison built from a loaded
    chart, however that chart was reached.

    `tostring` is typed for the standard library's `Element`, answering
    `str`, but openpyxl builds its tree with `lxml` at runtime, where the
    same call answers `bytes` -- the first `cast` says which of the two
    this reads. `from_tree` is typed to take the generic
    `_SerialisableTreeElement` protocol every `Serialisable` subclass
    shares and to answer `Self | None`; the element `fromstring` builds
    satisfies that protocol structurally, and a `<c:chartSpace>` root
    never parses to `None`, so the second and third `cast`s say what is
    true of this one call rather than of `Serialisable.from_tree` in
    general.
    """
    element = cast("_SerialisableTreeElement", fromstring(data))
    chart_space = cast("ChartSpace", ChartSpace.from_tree(element))
    return cast("bytes", tostring(chart_space.to_tree()))


def _chart_count(ws: Worksheet) -> int:
    """Return how many charts `ws` carries."""
    # `_charts` is undocumented and unstubbed, but it is the attribute
    # openpyxl's own reader and writer both use to hold a sheet's charts
    return len(ws._charts)  # type: ignore[attr-defined] # noqa: SLF001


def _chart_anchors(ws: Worksheet) -> list[tuple[int, int, int, int]]:
    """Return each of `ws`'s own charts' anchor cell, in the same order.

    `(col, row, colOff, rowOff)`, 0-based and read off `chart.anchor`'s
    own `_from` marker: a `OneCellAnchor`'s own position, the shape
    `ws.add_chart(chart, "H5")` writes and a round trip reads back as,
    rather than the string cell reference given at generation time.
    """
    anchors = []
    for chart in ws._charts:  # type: ignore[attr-defined] # noqa: SLF001
        marker = chart.anchor._from  # noqa: SLF001
        anchors.append((marker.col, marker.row, marker.colOff, marker.rowOff))
    return anchors


def _column_widths(ws: Worksheet) -> dict[str, float]:
    """Return every column of `ws` with an explicit width, keyed by letter.

    `ColumnDimension.width` is typed as a plain, never-`None` `float`,
    but a column `openpyxl` itself creates an entry for without this
    file ever setting a width carries `None` at runtime regardless --
    the stub states the type the descriptor validates a written value
    against, not what an unset one already holds.
    """
    return {
        letter: dim.width
        for letter, dim in ws.column_dimensions.items()
        if dim.width is not None  # type: ignore[redundant-expr]
    }


def _defined_names(container: Workbook | Worksheet) -> dict[str, str | None]:
    """Return `container`'s own defined names, keyed by name, valued by cell.

    A workbook-scoped name (`wb.defined_names`) and a sheet-scoped one
    (`ws.defined_names`) are both `DefinedNameDict`s of the same shape, so
    one function reads either -- btclib-org/bbt#134 is a sheet-scoped
    `prime` resolving to the wrong sheet's cell, which a comparison of
    cell values alone never reaches: every formula in every sheet reads
    `MOD(prime-A2,prime)` alike, whichever cell `prime` itself is bound
    to, so the drift is invisible until the binding itself is compared.
    """
    return {key: value.attr_text for key, value in container.defined_names.items()}


def _compare_sheet(  # noqa: PLR0913, PLR0917
    name: str,
    sheet: str,
    fresh_ws: Worksheet,
    committed_ws: Worksheet,
    fresh_charts: list[bytes],
    committed_charts: list[bytes],
) -> list[str]:
    """Return every way `committed_ws` departs from `fresh_ws`.

    `fresh_charts` and `committed_charts` are the whole workbook's own
    flat lists of already-serialised chart parts (`compare`'s own
    docstring on the order they share with `ws._charts`); this consumes
    `sheet`'s own share off the front of each, in place, so the next
    sheet finds its own share at the front in turn.
    """
    failures = []
    fresh_sheet_names = _defined_names(fresh_ws)
    committed_sheet_names = _defined_names(committed_ws)
    if fresh_sheet_names != committed_sheet_names:
        failures.append(
            f"  {name}:{sheet}: sheet defined names {committed_sheet_names} "
            f"committed, generate.py writes {fresh_sheet_names}",
        )
    fresh_cells, committed_cells = _cells(fresh_ws), _cells(committed_ws)
    for coordinate in sorted(set(fresh_cells) | set(committed_cells)):
        fresh_value = fresh_cells.get(coordinate)
        committed_value = committed_cells.get(coordinate)
        if fresh_value != committed_value:
            failures.append(
                f"  {name}:{sheet}!{coordinate}: generate.py writes "
                f"{fresh_value!r}, committed has {committed_value!r}",
            )

    fresh_widths, committed_widths = _column_widths(fresh_ws), _column_widths(
        committed_ws,
    )
    if fresh_widths != committed_widths:
        failures.append(
            f"  {name}:{sheet}: column widths {committed_widths} committed, "
            f"generate.py writes {fresh_widths}",
        )

    fresh_anchors, committed_anchors = _chart_anchors(fresh_ws), _chart_anchors(
        committed_ws,
    )
    if fresh_anchors != committed_anchors:
        failures.append(
            f"  {name}:{sheet}: chart anchors {committed_anchors} committed, "
            f"generate.py writes {fresh_anchors}",
        )

    fresh_count, committed_count = _chart_count(fresh_ws), _chart_count(committed_ws)
    sheet_fresh = fresh_charts[:fresh_count]
    sheet_committed = committed_charts[:committed_count]
    del fresh_charts[:fresh_count]
    del committed_charts[:committed_count]
    if fresh_count != committed_count:
        failures.append(
            f"  {name}:{sheet}: {committed_count} chart(s) committed, "
            f"generate.py writes {fresh_count}",
        )
        return failures
    for index, (fresh_chart, committed_chart) in enumerate(
        zip(sheet_fresh, sheet_committed, strict=True),
    ):
        if fresh_chart != committed_chart:
            failures.append(f"  {name}:{sheet}: chart {index} differs")
    return failures


def compare(
    name: str,
    fresh: Workbook,
    committed: Workbook,
    fresh_bytes: bytes,
    committed_bytes: bytes,
) -> list[str]:
    """Return every way `committed` departs from `fresh`, `name` prefixed."""
    failures = []
    fresh_names, committed_names = _defined_names(fresh), _defined_names(committed)
    if fresh_names != committed_names:
        failures.append(
            f"  {name}: workbook defined names {committed_names} committed, "
            f"generate.py writes {fresh_names}",
        )
    if fresh.sheetnames != committed.sheetnames:
        failures.append(
            f"  {name}: sheets {committed.sheetnames} committed, "
            f"generate.py writes {fresh.sheetnames}",
        )
        return failures

    # One flat list per side, in the same creation order `_chart_parts`
    # and `ws._charts` agree on; `_compare_sheet` consumes each sheet's
    # own count off the front of both, so a chart's own index into
    # either list never has to be recomputed from a sheet name.
    fresh_charts = [_chart_xml(part) for part in _chart_parts(fresh_bytes)]
    committed_charts = [_chart_xml(part) for part in _chart_parts(committed_bytes)]

    for sheet in fresh.sheetnames:
        failures.extend(
            _compare_sheet(
                name,
                sheet,
                fresh[sheet],
                committed[sheet],
                fresh_charts,
                committed_charts,
            ),
        )
    return failures


def main() -> int:
    """Regenerate every workbook, and compare it with what is committed."""
    generate = _load_generate()
    failures: list[str] = []
    workbooks: list[tuple[str, Workbook]] = [
        (spec.filename, generate.build_ec_workbook(spec)) for spec in generate.CURVES
    ]
    workbooks.append(
        (generate.FINITE_FIELDS_FILENAME, generate.build_finite_fields_workbook()),
    )
    if not workbooks:
        failures.append("  excel/: generate.py names no workbook, so nothing was read")
    for filename, built in workbooks:
        path = EXCEL_DIR / filename
        if not path.exists():
            failures.append(f"  excel/{filename}: generate.py names it, not committed")
            print(f"excel/{filename}: missing")
            continue
        fresh, fresh_bytes = _roundtrip(built)
        committed_bytes = path.read_bytes()
        committed = openpyxl.load_workbook(path)
        found = compare(
            f"excel/{filename}", fresh, committed, fresh_bytes, committed_bytes,
        )
        print(f"excel/{filename}: {'differs from' if found else 'matches'} generate.py")
        failures.extend(found)
    expected = {filename for filename, _ in workbooks}
    for path in sorted(EXCEL_DIR.glob("*.xlsx")):
        if path.name not in expected:
            failures.append(
                f"  excel/{path.name}: committed, generate.py does not name it",
            )
            print(f"excel/{path.name}: unexpected")
    for failure in failures:
        print(failure)
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
